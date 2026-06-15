import importlib.util
import json
from pathlib import Path
import sys
import tempfile
import unittest

import pandas as pd
from openpyxl import load_workbook


SCRIPTS = Path(__file__).parents[1] / "scripts"
sys.path.insert(0, str(SCRIPTS))


def load_script(name):
    path = SCRIPTS / name
    spec = importlib.util.spec_from_file_location(name.replace(".py", ""), path)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


CONTRACTS = load_script("stage_contracts.py")
STAGE0 = load_script("00_env_setup.py")
STAGE1 = load_script("01_data_prep.py")
STAGE2 = load_script("02_eda.py")
STAGE3 = load_script("03_feature_engineering.py")


class ContractTests(unittest.TestCase):
    def test_manifest_round_trip_and_required_formal_entry(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            config = root / "00_modeling_config.yaml"
            config.write_text(json.dumps({"model_type": "LR"}), encoding="utf-8")
            manifest_path = CONTRACTS.write_stage_manifest(
                root, 0, "completed", config, {"raw_data": "raw.csv"},
                {"config": str(config)}, [], 1,
            )
            manifest = CONTRACTS.load_previous_manifest(manifest_path, expected_stage=0)
            self.assertEqual(manifest["next_stage"], 1)
            self.assertTrue(Path(manifest["config_snapshot"]).exists())
            with self.assertRaises(ValueError):
                CONTRACTS.require_formal_entry(None, compatibility_mode=False, stage=1)
            next_manifest_path = CONTRACTS.write_stage_manifest(
                root, 1, "completed", config, {"previous_manifest": str(manifest_path)},
                {"train": str(root / "01_train.csv")}, [], 2,
            )
            next_manifest = json.loads(next_manifest_path.read_text(encoding="utf-8"))
            self.assertIn("config", next_manifest["outputs"])
            self.assertIn("train", next_manifest["outputs"])
            self.assertEqual(next_manifest["config_path"], str(config.resolve()))

    def test_output_workbook_contains_summary_inventory_and_pending(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            output = Path(temp_dir) / "00-output-list.xlsx"
            CONTRACTS.write_output_list(
                output,
                pd.DataFrame([{"status": "pending"}]),
                [Path(temp_dir) / "missing.csv"],
                pd.DataFrame([{"decision": "pending"}]),
            )
            workbook = load_workbook(output, read_only=True)
            self.assertEqual(
                set(workbook.sheetnames),
                {"summary", "file_inventory", "pending_decisions"},
            )
            workbook.close()


class Stage1Tests(unittest.TestCase):
    def test_stage0_confirmation_checklist_marks_blocking_items_pending(self):
        config = {
            "paths": {"raw_data_path": "raw.csv", "code_dir": "pending", "intermediate_dir": "out", "report_dir": "pending"},
            "fields": {"target_col": "label", "time_col": "pending", "sample_id_col": None, "sample_id_source_cols": []},
            "model_type": "pending",
            "enable_stage6_scoring": "pending",
            "psi_period_granularity": "pending",
            "sample_split": {"oot_method": "pending", "sample_method": "pending", "confirmed_by_user": False},
            "missing_definition": {"tokens": [], "blank_as_missing": True, "confirmed_by_user": False},
        }
        checklist = STAGE0.build_confirmation_checklist(config, "raw.csv")
        pending_items = set(checklist.loc[checklist["decision"].eq("pending"), "config_item"])
        self.assertIn("paths.code_dir", pending_items)
        self.assertIn("paths.report_dir", pending_items)
        self.assertIn("fields.time_col", pending_items)
        self.assertIn("model_type", pending_items)
        self.assertIn("sample_split.sample_method", pending_items)
        self.assertIn("missing_definition.confirmed_by_user", pending_items)

    def test_discover_sample_id_candidates_prefers_common_business_keys(self):
        columns = ["apply_no", "mobile_md5", "idcard_md5", "feature_a"]
        candidates = STAGE1.discover_sample_id_candidates(columns)
        self.assertEqual(candidates[:3], ["apply_no", "mobile_md5", "idcard_md5"])

    def test_stage1_rejects_unconfirmed_content_hash_fallback(self):
        config = {
            "fields": {"target_col": "label", "time_col": "app_dt", "sample_id_col": None, "sample_id_source_cols": []},
            "sample_split": {"oot_method": "months", "sample_method": "full", "confirmed_by_user": True},
            "missing_definition": {"tokens": [""], "blank_as_missing": True, "confirmed_by_user": True},
            "model_type": "LR",
        }
        with self.assertRaises(ValueError):
            STAGE1.validate_stage1_confirmations(config)

    def test_proportion_oot_keeps_same_timestamp_together(self):
        data = pd.DataFrame({
            "time": pd.to_datetime([
                "2024-01-01", "2024-01-02", "2024-01-03",
                "2024-01-03", "2024-01-04",
            ]),
            "target": [0, 0, 1, 0, 1],
        })
        modeling, oot = STAGE1.split_oot_by_proportion(data, "time", 0.4)
        self.assertEqual(set(oot["time"]), {pd.Timestamp("2024-01-03"), pd.Timestamp("2024-01-04")})
        self.assertFalse(set(modeling.index) & set(oot.index))

    def test_content_hash_sample_id_is_repeatable_and_distinguishes_duplicates(self):
        data = pd.DataFrame({"x": [1, 1, 2], "y": ["a", "a", "b"]})
        first = STAGE1.ensure_sample_id(data, None)
        second = STAGE1.ensure_sample_id(data, None)
        self.assertEqual(first["sample_id"].tolist(), second["sample_id"].tolist())
        self.assertEqual(first["sample_id"].nunique(), len(first))
        self.assertNotEqual(first.loc[0, "sample_id"], first.loc[1, "sample_id"])

    def test_composite_source_columns_drive_generated_id(self):
        data = pd.DataFrame({"customer": [1, 1], "application": [10, 11], "noise": ["a", "a"]})
        result = STAGE1.ensure_sample_id(data, None, ["customer", "application"])
        self.assertEqual(result["sample_id"].nunique(), 2)


class Stage3QualityActionTests(unittest.TestCase):
    def test_structured_quality_actions_apply_consistently(self):
        datasets = {
            "train": pd.DataFrame({"x": [-2, 1, 9], "cat": ["a", "b", "c"]}),
            "oot": pd.DataFrame({"x": [0, 20], "cat": ["a", "d"]}),
        }
        decisions = pd.DataFrame([
            {"feature": "x", "decision": "cap", "decision_detail": '{"lower": 0, "upper": 10}'},
            {"feature": "cat", "decision": "transform", "decision_detail": '{"action": "merge_categories", "mapping": {"a": "ab", "b": "ab"}}'},
        ])
        result, audit = STAGE3.apply_quality_actions(datasets, decisions)
        self.assertEqual(result["train"]["x"].tolist(), [0, 1, 9])
        self.assertEqual(result["oot"]["x"].tolist(), [0, 10])
        self.assertEqual(result["train"]["cat"].tolist(), ["ab", "ab", "c"])
        self.assertEqual(len(audit), 2)

    def test_unsupported_transform_is_blocked(self):
        decisions = pd.DataFrame([{
            "feature": "x", "decision": "transform",
            "decision_detail": '{"action": "log1p"}',
        }])
        with self.assertRaises(ValueError):
            STAGE3.apply_quality_actions({"train": pd.DataFrame({"x": [1]})}, decisions)


class Stage2ConfiguredQualityTests(unittest.TestCase):
    def test_optional_quality_checks_require_configuration(self):
        data = pd.DataFrame({"x": [None, None, 1], "cat": ["a", "a", "b"]})
        self.assertEqual(STAGE2.configured_quality_candidates(data, ["x", "cat"], {}), [])
        candidates = STAGE2.configured_quality_candidates(
            data, ["x", "cat"], {"high_missing_rate": 0.5, "sparse_category_rate": 0.4},
        )
        self.assertEqual({row["issue_type"] for row in candidates}, {"high_missing", "sparse_category"})


if __name__ == "__main__":
    unittest.main()
